"""Data export endpoints — CSV and XLSX export for metrics, suppliers, risk flags."""

import csv
import io
import logging
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from src.api.middleware.auth import require_auth
from src.db.database import fetch_metrics, fetch_risk_flags, fetch_suppliers

logger = logging.getLogger(__name__)

router = APIRouter()

ALLOWED_FORMATS = {"csv", "xlsx"}


def _require_format(format: str) -> str:
    if format not in ALLOWED_FORMATS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported format '{format}'. Allowed: {sorted(ALLOWED_FORMATS)}",
        )
    return format


# ---------------------------------------------------------------------------
# GET /api/export/metrics
# ---------------------------------------------------------------------------


@router.get("/metrics")
def export_metrics(
    format: str = Query("csv", description="Export format: csv or xlsx"),
    user: dict = Depends(require_auth),
) -> StreamingResponse:
    """Export metrics as CSV or XLSX. Org-scoped via auth token."""
    _require_format(format)
    org_id = user["org_id"]

    rows = fetch_metrics(org_id=org_id)
    if not rows:
        rows = []

    if format == "csv":
        return _metrics_csv(rows)
    else:
        return _metrics_xlsx(rows)


def _metrics_csv(rows: list[dict[str, Any]]) -> StreamingResponse:
    output = io.StringIO()
    fieldnames = [
        "id",
        "cluster",
        "value",
        "unit",
        "confidence",
        "source",
        "period",
        "recorded_at",
        "hash",
        "prev_hash",
    ]
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=metrics.csv"},
    )


def _metrics_xlsx(rows: list[dict[str, Any]]) -> StreamingResponse:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is not installed. Install with: pip install openpyxl",
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Metrics"

    headers = [
        "ID",
        "Cluster",
        "Value",
        "Unit",
        "Confidence",
        "Source",
        "Period",
        "Recorded At",
        "Hash",
        "Previous Hash",
    ]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font

    fieldnames = [
        "id",
        "cluster",
        "value",
        "unit",
        "confidence",
        "source",
        "period",
        "recorded_at",
        "hash",
        "prev_hash",
    ]
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, field in enumerate(fieldnames, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(field, ""))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=metrics.xlsx"},
    )


# ---------------------------------------------------------------------------
# GET /api/export/suppliers
# ---------------------------------------------------------------------------


@router.get("/suppliers")
def export_suppliers(
    format: str = Query("csv", description="Export format: csv or xlsx"),
    user: dict = Depends(require_auth),
) -> StreamingResponse:
    """Export suppliers as CSV or XLSX. Org-scoped via auth token."""
    _require_format(format)
    org_id = user["org_id"]

    rows = fetch_suppliers(org_id=org_id, limit=10000)
    if not rows:
        rows = []

    if format == "csv":
        return _suppliers_csv(rows)
    else:
        return _suppliers_xlsx(rows)


def _suppliers_csv(rows: list[dict[str, Any]]) -> StreamingResponse:
    output = io.StringIO()
    if not rows:
        writer = csv.writer(output)
        writer.writerow(
            [
                "ID",
                "Name",
                "Country",
                "Industry",
                "Tier",
                "Annual Spend USD",
                "Risk Tier",
                "Risk Score",
                "ESG Score",
                "Questionnaire Status",
                "Active Flags",
                "Certifications",
                "Created At",
            ]
        )
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=suppliers.csv"},
        )

    fieldnames = list(rows[0].keys())
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=suppliers.csv"},
    )


def _suppliers_xlsx(rows: list[dict[str, Any]]) -> StreamingResponse:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is not installed. Install with: pip install openpyxl",
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Suppliers"

    if not rows:
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            io.BytesIO(output.getvalue()),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=suppliers.xlsx"},
        )

    headers = list(rows[0].keys())
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, key in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))

    # Auto-size columns
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=suppliers.xlsx"},
    )


# ---------------------------------------------------------------------------
# GET /api/export/risk-flags
# ---------------------------------------------------------------------------


@router.get("/risk-flags")
def export_risk_flags(
    format: str = Query("csv", description="Export format: csv or xlsx"),
    user: dict = Depends(require_auth),
) -> StreamingResponse:
    """Export risk flags as CSV or XLSX. Org-scoped via auth token."""
    _require_format(format)
    org_id = user["org_id"]

    rows = fetch_risk_flags(org_id=org_id, limit=10000)
    if not rows:
        rows = []

    if format == "csv":
        return _risk_flags_csv(rows)
    else:
        return _risk_flags_xlsx(rows)


def _risk_flags_csv(rows: list[dict[str, Any]]) -> StreamingResponse:
    output = io.StringIO()
    if not rows:
        writer = csv.writer(output)
        writer.writerow(
            [
                "ID",
                "Flag Text",
                "Cluster",
                "Severity",
                "Days Overdue",
                "Priority Score",
                "Created At",
                "Acknowledged",
            ]
        )
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=risk_flags.csv"},
        )

    fieldnames = list(rows[0].keys())
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=risk_flags.csv"},
    )


def _risk_flags_xlsx(rows: list[dict[str, Any]]) -> StreamingResponse:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is not installed. Install with: pip install openpyxl",
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk Flags"

    if not rows:
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            io.BytesIO(output.getvalue()),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=risk_flags.xlsx"},
        )

    headers = list(rows[0].keys())
    header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, key in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=risk_flags.xlsx"},
    )
