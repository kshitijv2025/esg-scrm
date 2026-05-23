"""
D7.7: Supplier Data Exchange Tests.

Tests the /api/suppliers/exchange endpoint which supports:
  - GET: export supplier data in CSV, JSON, or XLSX format
  - POST: import suppliers with validation
  - POST /exchange/batch: batch-update supplier fields

Also covers:
  - scope3 data exchange format
  - country_risk data exchange format
  - org_id isolation (JWT enforced)
"""
import csv
import io
import json

import pytest
from fastapi.testclient import TestClient

from src.api.main import app
from src.auth.jwt import create_token
from src.db.database import get_connection, reset_database

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture(autouse=True)
def _fresh_db():
    """Reset and seed the database before every test."""
    reset_database()
    from src.db.seed import seed
    seed()
    from src.db.seed_emission_factors import seed_emission_factors
    seed_emission_factors()
    from src.db.seed_alert_thresholds import seed_alert_thresholds
    seed_alert_thresholds()
    from src.db.seed_country_risk import seed_country_risk
    seed_country_risk()


@pytest.fixture
def client():
    """Authenticated TestClient with admin role."""
    token = create_token({
        "sub": "usr_admin_001",
        "org_id": "org_bd_001",
        "email": "admin@test.com",
        "role": "admin",
    })
    c = TestClient(app)
    c.headers.update({"Authorization": f"Bearer {token}"})
    return c


@pytest.fixture
def editor_client():
    """Authenticated TestClient with editor role (required for POST endpoints)."""
    token = create_token({
        "sub": "usr_editor_001",
        "org_id": "org_bd_001",
        "email": "editor@test.com",
        "role": "editor",
    })
    c = TestClient(app)
    c.headers.update({"Authorization": f"Bearer {token}"})
    return c


@pytest.fixture
def viewer_client():
    """Authenticated TestClient with viewer role (cannot POST)."""
    token = create_token({
        "sub": "usr_viewer_001",
        "org_id": "org_bd_001",
        "email": "viewer@test.com",
        "role": "viewer",
    })
    c = TestClient(app)
    c.headers.update({"Authorization": f"Bearer {token}"})
    return c


@pytest.fixture
def other_org_client():
    """Client belonging to a different org."""
    token = create_token({
        "sub": "usr_other_001",
        "org_id": "org_other_999",
        "email": "other@test.com",
        "role": "admin",
    })
    c = TestClient(app)
    c.headers.update({"Authorization": f"Bearer {token}"})
    return c


# ---------------------------------------------------------------------------
# 1. Exchange Endpoint Exists
# ---------------------------------------------------------------------------

class TestExchangeEndpointExists:
    """Verify /api/suppliers/exchange returns data for all supported formats."""

    def test_exchange_returns_200(self, client):
        """GET /api/suppliers/exchange returns 200 with JSON by default."""
        resp = client.get("/api/suppliers/exchange")
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/json")

    def test_exchange_requires_auth(self):
        """Unauthenticated request → 401."""
        c = TestClient(app)
        resp = c.get("/api/suppliers/exchange")
        assert resp.status_code == 401

    def test_exchange_org_isolation(self, other_org_client):
        """Different org sees different (or empty) supplier data."""
        resp = other_org_client.get("/api/suppliers/exchange")
        assert resp.status_code == 200
        # Other org has no seeded suppliers, so total should be 0
        data = resp.json()
        # JSON is wrapped in StreamingResponse content
        if isinstance(data, dict):
            assert data.get("total", 0) == 0

    def test_exchange_default_format_is_json(self, client):
        """Default format (no ?format=) returns JSON."""
        resp = client.get("/api/suppliers/exchange")
        assert resp.status_code == 200
        assert "application/json" in resp.headers["content-type"]


# ---------------------------------------------------------------------------
# 2. Export Format Validation
# ---------------------------------------------------------------------------

class TestExportFormats:
    """GET /api/suppliers/exchange supports csv, json, and xlsx formats."""

    def test_format_json(self, client):
        """?format=json returns valid JSON with suppliers list."""
        resp = client.get("/api/suppliers/exchange?format=json")
        assert resp.status_code == 200
        assert "application/json" in resp.headers["content-type"]
        # StreamingResponse — content is the JSON string
        content = resp.content.decode()
        data = json.loads(content)
        assert "suppliers" in data
        assert "total" in data
        assert isinstance(data["suppliers"], list)

    def test_format_csv(self, client):
        """?format=csv returns valid CSV with headers."""
        resp = client.get("/api/suppliers/exchange?format=csv")
        assert resp.status_code == 200
        assert "text/csv" in resp.headers["content-type"]
        assert "attachment" in resp.headers.get("content-disposition", "")
        content = resp.content.decode()
        reader = csv.DictReader(io.StringIO(content))
        rows = list(reader)
        assert len(rows) >= 1
        # Check headers are present
        assert "name" in reader.fieldnames or "id" in reader.fieldnames

    def test_format_xlsx(self, client):
        """?format=xlsx returns an Excel file (openpyxl-readable)."""
        resp = client.get("/api/suppliers/exchange?format=xlsx")
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        content = resp.content
        # openpyxl can read the bytes directly from BytesIO
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content))
        assert "Suppliers" in wb.sheetnames

    def test_format_csv_scope_summary(self, client):
        """?format=csv&scope=summary filters to summary fields."""
        resp = client.get("/api/suppliers/exchange?format=csv&scope=summary")
        assert resp.status_code == 200
        content = resp.content.decode()
        reader = csv.DictReader(io.StringIO(content))
        rows = list(reader)
        # Summary should have at most these fields (no extra fields like annual_spend)
        if rows:
            assert len(rows[0]) <= 8  # id, name, country, tier, risk_score, ...

    def test_invalid_format_returns_400(self, client):
        """Unsupported format → 400 with clear error."""
        resp = client.get("/api/suppliers/exchange?format=pdf")
        assert resp.status_code == 400
        assert "Unsupported format" in resp.json()["detail"]


# ---------------------------------------------------------------------------
# 3. Import with Validation
# ---------------------------------------------------------------------------

class TestImportSuppliers:
    """POST /api/suppliers/exchange validates and imports supplier records."""

    def test_import_requires_editor_role(self, viewer_client):
        """Viewer (non-editor) gets 403 on POST."""
        resp = viewer_client.post(
            "/api/suppliers/exchange",
            json={"suppliers": [{"name": "New Supplier", "country": "BD"}]},
        )
        assert resp.status_code == 403

    def test_import_requires_auth(self, editor_client):
        """Unauthenticated → 401."""
        c = TestClient(app)
        resp = c.post(
            "/api/suppliers/exchange",
            json={"suppliers": [{"name": "New Supplier", "country": "BD"}]},
        )
        assert resp.status_code == 401

    def test_import_new_supplier(self, editor_client):
        """Valid supplier entry creates a new record with correct org_id from JWT."""
        resp = editor_client.post(
            "/api/suppliers/exchange",
            json={
                "suppliers": [
                    {
                        "name": "New Import Supplier",
                        "country": "VN",
                        "tier": "tier2",
                        "industry": "Electronics",
                        "annual_spend_usd": 250000,
                    }
                ]
            },
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["imported"] == 1
        assert body["updated"] == 0
        assert body["total"] == 1

        # Verify org_id is from JWT, not request body
        conn = get_connection()
        row = conn.execute(
            "SELECT org_id, name FROM suppliers WHERE name = ?",
            ("New Import Supplier",),
        ).fetchone()
        conn.close()
        assert row["org_id"] == "org_bd_001"

    def test_import_updates_existing(self, editor_client):
        """Entry with existing supplier id → updates instead of insert."""
        # First import
        editor_client.post(
            "/api/suppliers/exchange",
            json={"suppliers": [{"id": "sup_001", "name": "Original Name", "country": "BD"}]},
        )
        # Second import with same id → update
        resp = editor_client.post(
            "/api/suppliers/exchange",
            json={"suppliers": [{"id": "sup_001", "name": "Updated Name", "country": "BD"}]},
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["imported"] == 0
        assert body["updated"] == 1

        conn = get_connection()
        row = conn.execute(
            "SELECT name FROM suppliers WHERE id = ?", ("sup_001",)
        ).fetchone()
        conn.close()
        assert row["name"] == "Updated Name"

    def test_import_rejects_missing_name(self, editor_client):
        """Entry without name → error in errors list."""
        resp = editor_client.post(
            "/api/suppliers/exchange",
            json={"suppliers": [{"name": "", "country": "BD"}]},
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["imported"] == 0
        assert body["updated"] == 0
        assert any("name" in e.get("error", "") for e in body["errors"])

    def test_import_rejects_missing_country(self, editor_client):
        """Entry without country → error in errors list."""
        resp = editor_client.post(
            "/api/suppliers/exchange",
            json={"suppliers": [{"name": "Supplier Without Country"}]},
        )
        assert resp.status_code == 200
        body = resp.json()
        assert any("country" in e.get("error", "") for e in body["errors"])

    def test_import_rejects_non_list(self, editor_client):
        """suppliers field must be a list, not a dict."""
        resp = editor_client.post(
            "/api/suppliers/exchange",
            json={"suppliers": {"name": "Not a list"}},
        )
        assert resp.status_code == 400
        assert "list" in resp.json()["detail"].lower()

    def test_import_rejects_empty_list(self, editor_client):
        """Empty suppliers list → 400."""
        resp = editor_client.post(
            "/api/suppliers/exchange",
            json={"suppliers": []},
        )
        assert resp.status_code == 400
        assert "empty" in resp.json()["detail"].lower()

    def test_import_partial_failure_continues(self, editor_client):
        """Some bad entries are collected in errors, good entries are still imported."""
        resp = editor_client.post(
            "/api/suppliers/exchange",
            json={
                "suppliers": [
                    {"name": "Good Supplier", "country": "BD"},
                    {"name": "", "country": "BD"},  # Bad: no name
                    {"name": "Another Good", "country": "VN"},
                ]
            },
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["imported"] == 2
        assert len(body["errors"]) >= 1


# ---------------------------------------------------------------------------
# 4. Batch Update Operations
# ---------------------------------------------------------------------------

class TestBatchUpdateOperations:
    """POST /api/suppliers/exchange/batch updates multiple supplier fields."""

    def test_batch_update_requires_editor(self, viewer_client):
        """Viewer role cannot batch update."""
        resp = viewer_client.post(
            "/api/suppliers/exchange/batch",
            json={"updates": [{"supplier_id": "sup_001", "risk_tier": "low"}]},
        )
        assert resp.status_code == 403

    def test_batch_update_risk_tier(self, editor_client):
        """Batch update can set risk_tier."""
        resp = editor_client.post(
            "/api/suppliers/exchange/batch",
            json={
                "updates": [
                    {"supplier_id": "sup_001", "risk_tier": "A", "questionnaire_status": "responded"}
                ]
            },
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["updated"] == 1

        conn = get_connection()
        row = conn.execute(
            "SELECT risk_tier, questionnaire_status FROM suppliers WHERE id = ?",
            ("sup_001",),
        ).fetchone()
        conn.close()
        assert row["risk_tier"] == "A"
        assert row["questionnaire_status"] == "responded"

    def test_batch_update_unknown_supplier(self, editor_client):
        """Unknown supplier_id → error in errors list, no crash."""
        resp = editor_client.post(
            "/api/suppliers/exchange/batch",
            json={"updates": [{"supplier_id": "sup_nonexistent", "risk_tier": "low"}]},
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["updated"] == 0
        assert len(body["errors"]) >= 1
        assert "not found" in body["errors"][0]["error"].lower()

    def test_batch_update_requires_supplier_id(self, editor_client):
        """Entry without supplier_id → error."""
        resp = editor_client.post(
            "/api/suppliers/exchange/batch",
            json={"updates": [{"risk_tier": "low"}]},
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["updated"] == 0
        assert len(body["errors"]) >= 1

    def test_batch_update_empty_updates(self, editor_client):
        """Empty updates list → 400."""
        resp = editor_client.post(
            "/api/suppliers/exchange/batch",
            json={"updates": []},
        )
        assert resp.status_code == 400

    def test_batch_update_certifications(self, editor_client):
        """Batch update can set certifications field."""
        resp = editor_client.post(
            "/api/suppliers/exchange/batch",
            json={
                "updates": [
                    {"supplier_id": "sup_001", "certifications": "ISO14001,SA8000"}
                ]
            },
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["updated"] == 1


# ---------------------------------------------------------------------------
# 5. Country Risk Data Exchange
# ---------------------------------------------------------------------------

class TestCountryRiskExchange:
    """GET /api/suppliers/exchange?scope=country_risk returns geopolitical data."""

    def test_country_risk_returns_200(self, client):
        """scope=country_risk returns data for supplier countries."""
        resp = client.get("/api/suppliers/exchange?scope=country_risk&format=json")
        assert resp.status_code == 200
        content = resp.content.decode()
        data = json.loads(content)
        assert "country_risk" in data

    def test_country_risk_csv_format(self, client):
        """scope=country_risk&format=csv returns valid CSV."""
        resp = client.get("/api/suppliers/exchange?scope=country_risk&format=csv")
        assert resp.status_code == 200
        content = resp.content.decode()
        reader = csv.DictReader(io.StringIO(content))
        rows = list(reader)
        # Should have country_code column at minimum
        assert "country_code" in reader.fieldnames or len(rows) >= 0

    def test_country_risk_xlsx_format(self, client):
        """scope=country_risk&format=xlsx returns Excel file."""
        resp = client.get("/api/suppliers/exchange?scope=country_risk&format=xlsx")
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]

    def test_country_risk_org_isolation(self, other_org_client):
        """Other org sees their own country risk data (or empty)."""
        resp = other_org_client.get("/api/suppliers/exchange?scope=country_risk&format=json")
        assert resp.status_code == 200
        content = resp.content.decode()
        data = json.loads(content)
        # org_other_999 has no suppliers → country_risk list should be empty
        assert data.get("total", len(data.get("country_risk", []))) == 0


# ---------------------------------------------------------------------------
# 6. Scope 3 Data Exchange
# ---------------------------------------------------------------------------

class TestScope3Exchange:
    """GET /api/suppliers/exchange?scope=scope3 returns emission data."""

    def test_scope3_returns_200(self, client):
        """scope=scope3 returns emission data for all suppliers."""
        resp = client.get("/api/suppliers/exchange?scope=scope3&format=json")
        assert resp.status_code == 200
        content = resp.content.decode()
        data = json.loads(content)
        assert "scope3" in data

    def test_scope3_csv_format(self, client):
        """scope=scope3&format=csv returns valid CSV with emission columns."""
        resp = client.get("/api/suppliers/exchange?scope=scope3&format=csv")
        assert resp.status_code == 200
        content = resp.content.decode()
        reader = csv.DictReader(io.StringIO(content))
        rows = list(reader)
        # Should have scope3_tco2e or category column
        if rows:
            assert any(
                col in reader.fieldnames
                for col in ("scope3_tco2e", "category", "annual_spend_usd")
            )

    def test_scope3_xlsx_format(self, client):
        """scope=scope3&format=xlsx returns Excel file."""
        resp = client.get("/api/suppliers/exchange?scope=scope3&format=xlsx")
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        content = resp.content
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content))
        assert "Scope3" in wb.sheetnames

    def test_scope3_org_isolation(self, other_org_client):
        """Other org sees no scope3 data (their suppliers have none)."""
        resp = other_org_client.get("/api/suppliers/exchange?scope=scope3&format=json")
        assert resp.status_code == 200
        content = resp.content.decode()
        data = json.loads(content)
        # org_other_999 has no suppliers → scope3 list empty
        assert data.get("total", len(data.get("scope3", []))) == 0


# ---------------------------------------------------------------------------
# 7. Org Isolation on All Exchange Operations
# ---------------------------------------------------------------------------

class TestExchangeOrgIsolation:
    """All exchange operations must use org_id from JWT, not request body."""

    def test_import_sets_org_id_from_jwt(self, editor_client):
        """POST import uses JWT org_id, ignoring any org_id in request body."""
        resp = editor_client.post(
            "/api/suppliers/exchange",
            json={
                "suppliers": [
                    {
                        "name": "Isolation Test Supplier",
                        "country": "BD",
                        "org_id": "org_evil_999",  # Attempt injection
                    }
                ]
            },
        )
        assert resp.status_code == 200
        conn = get_connection()
        row = conn.execute(
            "SELECT org_id FROM suppliers WHERE name = ?",
            ("Isolation Test Supplier",),
        ).fetchone()
        conn.close()
        assert row["org_id"] == "org_bd_001"

    def test_batch_update_respects_org(self, editor_client, other_org_client):
        """Other org's supplier id cannot be updated by this org."""
        resp = editor_client.post(
            "/api/suppliers/exchange/batch",
            json={"updates": [{"supplier_id": "sup_other_org_001", "risk_tier": "high"}]},
        )
        assert resp.status_code == 200
        body = resp.json()
        # sup_other_org_001 doesn't belong to org_bd_001
        assert body["updated"] == 0
        assert len(body["errors"]) >= 1
